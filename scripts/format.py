from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border

def main(file, start_row, end_row):
    """
    Receives literal Excel row numbers from app.py — no offset applied here.
    app.py is responsible for all row number conversions.
    """
    EXCEL_FILE = file
    COLORS = {
        "good":       "C6EFCE",
        "no_results": "FCE4D6",
        "no_author":  "FFF2CC",
        "bad_search": "F4CCCC",
        "error":      "EA9999",
        "na":         "D9D9D9",
        "other":      "FFE599",
    }
    NO_FILL   = PatternFill(fill_type=None)
    NO_BORDER = Border()

    def get_color(value):
        if value is None:
            return None
        v = str(value).strip()
        try:
            float(v)
            return COLORS["good"]
        except ValueError:
            pass
        if v.isdigit() or v == "eb" or v.endswith("eb"):
            return COLORS["good"]
        elif v == "NO RESULTS":
            return COLORS["no_results"]
        elif v == "NO AUTHOR":
            return COLORS["no_author"]
        elif v == "BAD SEARCH":
            return COLORS["bad_search"]
        elif v == "ERROR":
            return COLORS["error"]
        elif v == "N/A":
            return COLORS["na"]
        else:
            return COLORS["other"]

    wb = load_workbook(EXCEL_FILE)

    for ws in wb.worksheets:
        for table in list(ws.tables.values()):
            del ws.tables[table.displayName]

    ws = wb.active

    ol_col = None
    for cell in ws[1]:
        if cell.value == "OL":
            ol_col = cell.column
            break

    if ol_col is None:
        print("[FORMAT] Could not find OL column — skipping formatting.")
        return

    print(f"[FORMAT] Formatting Excel rows {start_row} to {end_row}...")

    for row in ws.iter_rows(min_row=start_row, max_row=end_row):
        for cell in row:
            cell.fill   = NO_FILL
            cell.border = NO_BORDER
        ol_cell   = row[ol_col - 1]
        hex_color = get_color(ol_cell.value)
        if hex_color:
            ol_cell.fill = PatternFill(patternType="solid", fgColor=hex_color)

    wb.save(EXCEL_FILE)
    print("[FORMAT] Done.")

if __name__ == "__main__":
    print("Not a solo file — run through app.py")