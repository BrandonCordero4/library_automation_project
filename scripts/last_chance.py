import pandas as pd
import re
from difflib import SequenceMatcher
from urllib.parse import quote

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from unidecode import unidecode

# ============================================================
# CONFIG
# ============================================================

EXCEL_FILE = "../books/books_test.xlsx"

IGNORE_WORDS = {
    "by", "By", "Edited", "edited", "editor", "Editor", "editors", "Editors",
    "Sir", "Baron", "Lord", "Lady", "Dr", "Rev", "Professor", "Prof",
    "Jr", "Sr", "II", "III", "IV", "Illus", "illus", "illustrated", "Illustrated"
}

# ============================================================
# WAIT HELPERS
# ============================================================

def wait_for_detail_panel(page):
    print("  [WAIT] Waiting for detail panel to appear...")
    page.wait_for_selector(
        "//span[@data-details-label='title']",
        timeout=20000
    )
    page.wait_for_function(
        """() => {
            const el = document.evaluate(
                "//span[@data-details-label='title']/../../div[contains(@class,'item-details-element-container')]",
                document, null, XPathResult.FIRST_ORDERED_NODE_TYPE, null
            ).singleNodeValue;
            return el && el.innerText.trim().length > 0;
        }""",
        timeout=20000
    )
    print("  [WAIT] Detail panel ready.")


def wait_for_schools(page):
    print("  [WAIT] Waiting for school list to settle...")
    page.wait_for_timeout(300)
    count = page.locator("p.inst-name").count()
    print(f"  [WAIT] School list ready. Count: {count}")
    return count


# ============================================================
# HELPERS  (identical to mark9.py)
# ============================================================

def is_good_result(result):
    if isinstance(result, (int, float)):
        return True
    if isinstance(result, str):
        if result in {"NO AUTHOR", "N/A", "NO RESULTS", "BAD SEARCH"}:
            return True
        if result == "eb" or result.endswith("eb") or result.isdigit():
            return True
    return False


def save_to_sheet(row_number, value):
    global df
    df.at[row_number, "OL"] = value

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    ol_col = None
    for cell in ws[1]:
        if cell.value == "OL":
            ol_col = cell.column
            break

    ws.cell(row=row_number + 2, column=ol_col).value = value
    wb.save(EXCEL_FILE)

    if not is_good_result(value):
        print(f"  [SAVE] Flagged for review: ROW {row_number + 2} = '{value}'")


def get_current_record_id(page):
    try:
        return page.locator(
            "div.list-item-primary-content.result-item-primary-content[id*='FULL_VIEW']"
        ).first.get_attribute("data-recordid")
    except Exception:
        return None


def click_down_and_wait(page):
    down_button = page.locator('md-icon[md-svg-icon="primo-ui:chevron-down"]')
    count = down_button.count()
    print(f"  [CHEVRON] Chevron button found: {count > 0}")
    if not count:
        print("  [CHEVRON] No chevron — end of records.")
        return False

    current_id = get_current_record_id(page)
    print(f"  [CHEVRON] Clicking. Current record ID: {current_id}")
    down_button.click(timeout=4000, force=True)
    print("  [CHEVRON] Click fired. Waiting for record ID to change...")

    try:
        page.wait_for_function(
            """(prev_id) => {
                const el = document.querySelector(
                    "div.list-item-primary-content.result-item-primary-content[id*='FULL_VIEW']"
                );
                return el && el.getAttribute('data-recordid') !== prev_id;
            }""",
            arg=current_id,
            timeout=5000
        )
        new_id = get_current_record_id(page)
        print(f"  [CHEVRON] Record ID changed: {current_id} → {new_id}")
        return True
    except Exception:
        print(f"  [CHEVRON] Record ID did not change — end of records.")
        return False


def split_initials(names):
    fixed = []
    for name in names:
        if len(name) > 1 and name.isupper():
            fixed.extend(list(name))
        else:
            fixed.append(name)
    return fixed


def make_result(book_count, contains_ebook):
    if contains_ebook and book_count == 0:
        return "eb"
    if contains_ebook:
        return f"{book_count}eb"
    return book_count


def has_ebook(page):
    try:
        record_id = get_current_record_id(page)
    except Exception:
        print("  [EBOOK] Could not get record ID")
        return False

    if not record_id:
        print("  [EBOOK] No record ID found")
        return False

    found = page.locator(
        f"button[aria-labelledby^='{record_id}'] span.availability-status.not_restricted"
    ).count() > 0

    print(f"  [EBOOK] Record ID: {record_id} | Ebook detected: {found}")
    return found


def clean_title(text):
    text = text.lower()
    text = re.sub(r"[^\w\s]", " ", text)
    return " ".join(text.split())


def exact_title_match(sheet_title, web_title):
    def smash(t):
        return re.sub(r"[^a-z0-9]", "", t.lower())
    return smash(sheet_title) in smash(web_title)


def author_filter(author):
    if not author:
        return ""
    author = re.sub(r"[.,;:/()\[\]]", " ", author)
    words = author.split()
    names = [
        w for w in words
        if (len(w) == 1 and w.isalpha()) or (w and w[0].isupper())
    ]
    return " ".join(names)


def extract_author_names(author_text):
    author_names = [
        w for w in author_text.split()
        if w and w[0].isupper() and w not in IGNORE_WORDS
    ]
    author_names = split_initials(author_names)

    fixed_names = []
    i = 0
    while i < len(author_names):
        if i + 1 < len(author_names) and author_names[i] in {"De", "Van", "Von", "Mac", "Mc"}:
            fixed_names.append(author_names[i] + author_names[i + 1])
            i += 2
        else:
            fixed_names.append(author_names[i])
            i += 1
    return fixed_names


def title_matches(sheet_title, web_title):
    sheet_words = set(clean_title(sheet_title).split())
    web_words   = set(clean_title(web_title).split())
    if not sheet_words or not web_words:
        return False
    overlap  = len(sheet_words & web_words)
    required = min(len(sheet_words), len(web_words))
    return overlap >= 0.8 * required


def author_matches(author_names, web_author_names):
    unused_web_names   = web_author_names.copy()
    missing_full_names = []
    missing_initials   = []
    matched_initials   = 0
    matched_full_names = 0

    for name in author_names:
        if len(name) == 1:
            continue
        found = False
        if name in unused_web_names:
            unused_web_names.remove(name)
            found = True
        else:
            parts = name.split("-")
            if len(parts) > 1:
                if sum(1 for p in parts if p in unused_web_names) > 0:
                    for p in parts:
                        if p in unused_web_names:
                            unused_web_names.remove(p)
                    found = True
            if not found:
                for web_name in unused_web_names:
                    if len(web_name) == 1 and name.startswith(web_name):
                        unused_web_names.remove(web_name)
                        found = True
                        break
        if found:
            matched_full_names += 1
        else:
            missing_full_names.append(name)

    for name in author_names:
        if len(name) != 1:
            continue
        for web_name in unused_web_names:
            if web_name.startswith(name):
                unused_web_names.remove(web_name)
                matched_initials += 1
                break
        else:
            missing_initials.append(name)

    if len(missing_full_names) == 1 and len(unused_web_names) == 1:
        similarity = SequenceMatcher(
            None,
            unidecode(missing_full_names[0]).lower(),
            unidecode(unused_web_names[0]).lower()
        ).ratio()
        print(f"  [FUZZY] {missing_full_names[0]} vs {unused_web_names[0]} = {similarity:.2%}")
        if similarity >= 0.85:
            matched_full_names += 1
            missing_full_names.clear()
            unused_web_names.clear()

    correct_author = matched_full_names >= 1 and matched_full_names + matched_initials >= 2
    return {
        "correct":            correct_author,
        "matched_full_names": matched_full_names,
        "missing_full_names": missing_full_names,
        "missing_initials":   missing_initials,
        "web_author_names":   web_author_names,
    }


def get_web_title(page):
    return (
        page.locator("//span[@data-details-label='title']")
        .locator("xpath=../../div[contains(@class,'item-details-element-container')]")
        .inner_text()
        .strip()
    )


def correct_book(page, title, author):
    raw_web_title  = get_web_title(page)
    raw_web_author = (
        page.locator("//span[@data-details-label='creator']")
        .locator("xpath=../../div[contains(@class,'item-details-element-container')]")
        .inner_text()
        .strip()
    )

    web_author    = author_filter(raw_web_author)
    correct_title = title_matches(title, raw_web_title)

    print(f"  [CHECK] WEB TITLE:  [{raw_web_title}]")
    print(f"  [CHECK] WEB AUTHOR: [{web_author}]")
    print(f"  [CHECK] Title Match: {correct_title}")

    author_names     = extract_author_names(author)
    web_author_names = extract_author_names(web_author)
    author_result    = author_matches(author_names, web_author_names)

    print(f"  [CHECK] Matched Full Names: {author_result['matched_full_names']}")
    print(f"  [CHECK] Author Match: {author_result['correct']}")

    if not author_result["correct"]:
        print(f"  [CHECK] !! Author names:     {author_names}")
        print(f"  [CHECK] !! Web author names: {author_result['web_author_names']}")
        print(f"  [CHECK] !! Missing full:     {author_result['missing_full_names']}")
        print(f"  [CHECK] !! Missing initials: {author_result['missing_initials']}")

    return {
        "correct":        correct_title and author_result["correct"],
        "web_title":      raw_web_title,
        "web_author":     web_author,
        "correct_title":  correct_title,
        "correct_author": author_result["correct"],
    }


def is_correct_record(page, title, author, no_author_mode):
    wait_for_detail_panel(page)
    if no_author_mode:
        web_title = get_web_title(page)
        passed    = exact_title_match(title, web_title)
        print(f"  [CHECK] WEB TITLE (no-author): [{web_title}]")
        print(f"  [CHECK] Exact Title Match: {passed}")
        return passed, web_title
    else:
        result = correct_book(page, title, author)
        return result["correct"], result["web_title"]


def count_holdings(page, title, author, no_author_mode):
    book_count     = 0
    book_number    = 1
    contains_ebook = False

    while True:
        print(f"\n  --- Record {book_number} ---")

        correct, web_title = is_correct_record(page, title, author, no_author_mode)

        if not correct:
            print(f"\n  [LOOP] Stopped — wrong book: [{web_title}]")
            result = make_result(book_count, contains_ebook)
            save_to_sheet(row_number, result)
            print(f"  [SAVE] Saved: '{result}'")
            break

        if has_ebook(page):
            contains_ebook = True

        schools = wait_for_schools(page)
        book_count += schools
        print(f"  [COUNT] Book {book_number:<2} | Schools: {schools:<3} | Running Total: {book_count}")
        print("  " + "-" * 48)
        book_number += 1

        if not click_down_and_wait(page):
            result = make_result(book_count, contains_ebook)
            print("  [LOOP] Reached end of records.")
            save_to_sheet(row_number, result)
            print(f"  [SAVE] Saved: '{result}'")
            break

# ============================================================
# MAIN — cleanup: retry only ERROR and HOW? rows
# ============================================================
def main():
    df = pd.read_excel(EXCEL_FILE)
    df["OL"] = df["OL"].astype("object")

    # Find every row that needs a retry
    retry_rows = [
        i for i, val in enumerate(df["OL"])
        if isinstance(val, str) and (val == "ERROR" or val.startswith("HOW?"))
    ]

    print(f"Found {len(retry_rows)} rows to retry: {[r + 2 for r in retry_rows]}")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        page    = browser.new_page()

        for row_number in retry_rows:
            print("\n\n\n" + "=" * 60)
            print(f"ROW {row_number + 2}  (retrying '{df['OL'].iloc[row_number]}')")
            print("=" * 60)

            try:
                parts     = str(df["TITLE"].iloc[row_number]).strip().split("/", 1)
                raw_title = parts[0].strip()

                search_title = unidecode(raw_title)
                search_title = re.sub(r";.*", "", search_title).strip()
                search_title = " ".join(search_title.split())
                title        = raw_title

                raw_author = parts[1].strip() if len(parts) > 1 else "NONE"
                author     = " ".join(re.sub(r"\band\b", " ", author_filter(raw_author), flags=re.IGNORECASE).split())
                author     = " ".join(w for w in author.split() if w not in IGNORE_WORDS)

                print(f"SEARCH TITLE : [{title}]")
                print(f"SEARCH AUTHOR: [{author}]")

                if not title or title.lower() == "nan":
                    print("[MAIN] Empty title → N/A")
                    save_to_sheet(row_number, "N/A")
                    continue

                no_author = raw_author == "NONE" or not author

                if no_author and len(title.split()) < 4:
                    print("[MAIN] Title too short for no-author search → NO AUTHOR")
                    save_to_sheet(row_number, "NO AUTHOR")
                    continue

                encoded_title  = quote(search_title)
                encoded_author = quote(author)

                if no_author:
                    print("[MAIN] No author — title-only mode.")
                    url = (
                        "https://primo.denison.edu/discovery/search?"
                        f"query=title,contains,{encoded_title},AND"
                        "&pfilter=rtype,exact,books,AND"
                        "&tab=New_Network_Catalog"
                        "&search_scope=New_Discovery_Network"
                        "&vid=01OHIOLINK_DENISON:DENISON&lang=en&mode=advanced&offset=0"
                    )
                else:
                    url = (
                        "https://primo.denison.edu/discovery/search?"
                        f"query=title,contains,{encoded_title},AND"
                        f"&query=creator,contains,{encoded_author},AND"
                        "&pfilter=rtype,exact,books,AND"
                        "&tab=New_Network_Catalog"
                        "&search_scope=New_Discovery_Network"
                        "&vid=01OHIOLINK_DENISON:DENISON&lang=en&mode=advanced&offset=0"
                    )

                print(f"[MAIN] Navigating to search URL...")
                page.goto(url)

                print("[MAIN] Waiting for results or no-records message...")
                page.wait_for_function(
                    """() =>
                        document.querySelectorAll('mark').length > 0 ||
                        document.body.innerText.includes('No records found')
                    """,
                    timeout=10000
                )
                print("[MAIN] Page content ready.")

                no_results = page.get_by_text("No records found").count()
                if no_results and not no_author:
                    print("[MAIN] No results. Retrying with search button...")
                    try:
                        page.locator('button[data-qa="advanced-search-btn"]').click()
                        page.wait_for_function(
                            """() =>
                                document.querySelectorAll('mark').length > 0 ||
                                document.body.innerText.includes('No records found')
                            """,
                            timeout=10000
                        )
                        no_results = page.get_by_text("No records found").count()
                    except Exception as e:
                        print(f"[MAIN] Retry failed: {e}")

                if no_results:
                    print("[MAIN] No results → NO RESULTS")
                    save_to_sheet(row_number, "NO RESULTS")
                    continue

                top_result = page.locator("mark")
                if not top_result.count():
                    print("[MAIN] No mark element found → HOW? Check 'mark'")
                    save_to_sheet(row_number, "HOW? Check 'mark'")
                    continue

                print(f"[MAIN] Clicking top result...")
                top_result.nth(0).click()

                wait_for_detail_panel(page)
                print("[MAIN] Detail panel loaded after clicking top result.")

                correct, _ = is_correct_record(page, title, author, no_author)
                if not correct:
                    label = "NO AUTHOR" if no_author else "BAD SEARCH"
                    print(f"[MAIN] First result wrong → {label}")
                    save_to_sheet(row_number, label)
                    continue

                print("[MAIN] First result confirmed. Starting holdings count.")
                count_holdings(page, title, author, no_author)

            except Exception as e:
                print(f"[MAIN] ERROR: {e}")
                save_to_sheet(row_number, "ERROR")
                print(f"  [SAVE] Flagged for review: ROW {row_number + 2}")

        input("\nDone. Press ENTER to close")
        
        browser.close()

if __name__ == "__main__":
    main()